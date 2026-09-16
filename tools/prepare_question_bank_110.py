import argparse
import re
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "829题库.xlsx"
MANUAL_SHEET = "逐题人工挖空"

MNEMONICS = {
    "性蜜罐漏拔稳",
    "高侵交（胶）体（替）毛（冒）",
    "膨固破沉锁",
    "密桥类完活",
    "热砂通渗侧隔低裂",
    "安稳护衡调配眼影粒腐",
    "速步术式原层",
    "地位独规",
    "天性关压",
    "基发适提",
    "厚隔技物",
    "术能环稳效速率",
    "则透后速造",
    "速调一面",
    "要原储层工技输流经地温驱指（可数）",
    "描藏油地经综",
    "程试监预",
    "分划同适剔细",
    "整密稀换局",
    "强粒杂便密",
    "地能转错情",
    "悬砂滤网摩擦稳定，这是低配的经济。",
}

REPLACEMENTS = {
    "配置而成": "配制而成",
    "配置高密度钻井液": "配制高密度钻井液",
    "普通射孔完共": "普通射孔完井",
    "转动转量很困难": "转动钻柱很困难",
    "井简中的": "井筒中的",
    "带出井简": "带出井筒",
    "井简内": "井筒内",
    "钻县": "钻具",
    "螺杆钻县": "螺杆钻具",
    "完并是沟通": "完井是沟通",
    "泡沫完并液": "泡沫完井液",
    "向并内": "向井内",
    "携带赃物": "携带脏物",
    "气审井": "气窜井",
    "关闭气窜并": "关闭气窜井",
    "分层开采工艺地作用": "分层开采工艺的作用",
    "通过并网": "通过井网",
    "两口并同时渗流": "两口井同时渗流",
    "开发并是指": "开发井是指",
    "本并低产": "本井低产",
    "裸漏": "裸露",
    "保待较长稳产期": "保持较长稳产期",
    "地层破裂压裂的变化": "地层破裂压力的变化",
    "有效作用距和裂缝": "有效作用距离和裂缝",
    "解堵化。": "解堵酸化。",
    "孔隙迁曲度": "孔隙迂曲度",
    "油单向流动区": "油单相流动区",
    "残余油饱和度S。": "残余油饱和度 Sor",
    "残余油饱和度Sor": "残余油饱和度 Sor",
    "束缚水饱和度Swi": "束缚水饱和度 Swi",
    "井底流压pwf": "井底流压 pwf",
    "CaMg(CO)3": "CaMg(CO₃)₂",
    "CaCO3": "CaCO₃",
    "粘土": "黏土",
    "粘度": "黏度",
    "粘稠": "黏稠",
    "浴井解卡": "泡油解卡",
    "钻井液油管": "洗井液由油管",
    "洗井液套管": "洗井液由套管",
    "起起升作用": "起升作用",
    "携形筒": "楔形筒",
    "油梁": "游梁",
    "相接处理": "相继进行",
    "顿钻钻井": "冲击钻井",
    "气压降大小": "其压降大小",
}

TITLE_OVERRIDES = {
    12: "束缚水的定义",
    18: "地层油溶解气油比随压力变化的关系",
    19: "地层油的体积系数 Bo",
    21: "地下原油体积的影响因素",
    27: "天然气的体积系数 Bg",
    32: "压力系数的作用",
    36: "地温梯度",
    48: "钻机八大系统（一）",
    49: "钻机八大系统（二）",
    51: "常用钻井工具（一）",
    52: "常用钻井工具（二）",
    55: "钻井液的分类",
    59: "气体型钻井流体",
    60: "合成基钻井液",
    62: "钻井液密度",
    66: "钻井液流变性",
    67: "宾汉模式与幂律模式的参数",
    73: "钻井液 pH 的控制范围及作用",
    74: "钻井液碱度",
    76: "起升系统的组成",
    77: "钻井液循环系统的组成",
    78: "地面旋转系统的组成",
    79: "动力驱动系统",
    81: "控制系统",
    84: "钻铤的位置",
    94: "钻井液的功用（24年考过）",
    112: "井喷的处理",
    113: "固井技术",
    126: "套管强度",
    136: "油井水泥的 API 级别",
    140: "水泥浆性能",
    141: "水泥浆的稠化时间",
    153: "裸眼完井法",
    154: "射孔完井法",
    156: "尾管射孔完井",
    158: "割缝衬管完井法",
    160: "后期割缝衬管完井",
    168: "套管头",
    183: "地层压力、井底流动压力和井口压力",
    185: "完井液的功能要求（简答）",
    195: "通井方式",
    202: "正、反循环洗井",
    210: "气举排液",
    228: "试油",
    230: "生产试验区",
    231: "基础井网",
    242: "开发层系与分层开采工艺的关系",
    244: "高、低渗透层合采的影响",
    246: "天然能量",
    250: "油藏驱动方式",
    256: "气压驱动",
    263: "渗流速度",
    264: "渗流形式",
    266: "稳定流动",
    268: "不稳定流动",
    277: "注水方式类型",
    312: "油井生产系统中的流动过程",
    323: "普通节点与函数节点",
    352: "基质酸化",
    354: "压裂液的作用",
    355: "压裂液按注入阶段的分类",
    358: "支撑剂",
    368: "常规酸化与酸压",
}

# Text on the first line is an answer-bearing statement, not a usable prompt.
# Keep it below the concise prompt so none of the user's manual blanks are lost.
TITLE_ONLY_ROWS = {18, 27, 48, 49, 51, 52, 59, 60, 94, 156, 160, 185}
ANSWER_BEARING_TITLES = set(TITLE_OVERRIDES) - TITLE_ONLY_ROWS - {19}

QUESTION_TYPES = {
    "definition": "定义",
    "classification": "分类",
    "condition": "条件",
    "mechanism": "机理",
    "effect": "作用",
    "result": "结果",
    "formula": "公式",
    "value": "数值",
    "mnemonic": "口诀",
    "list": "条目",
    "example": "例子",
}


def normalize_punctuation(text: str) -> str:
    text = str(text or "").replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\u00a0", " ").replace("\u3000", " ")
    text = text.replace(",", "，").replace(";", "；")
    text = re.sub(r"(?<=\d)~(?=\d)", "～", text)
    text = re.sub(r"(?<=\d)-(?=\d)", "～", text)
    text = text.replace("•", "·")
    text = re.sub(r"(?m)^\s*[（(](\d+)[）)]\s*", lambda m: f"（{m.group(1)}）", text)
    text = re.sub(r"(?m)^\s*(\d+)[）)]\s*", lambda m: f"（{m.group(1)}）", text)
    # Chinese prose uses full-width punctuation; mathematical parentheses are
    # left intact because replacing them would change formula readability.
    text = re.sub(r"[：:]\s*$", "", text)
    return text


def normalize_units(text: str) -> str:
    replacements = {
        "1mPa·s(毫帕秒)=1cP(厘泊)": "1 mPa·s（毫帕秒）=1 cP（厘泊）",
        "1mPa·s=1cP": "1 mPa·s=1 cP",
        "1P(泊)=1000cP": "1 P（泊）=100 cP",
        "1 P(泊)=1000 cP": "1 P（泊）=100 cP",
        "1Pa·s=1000mP·s": "1 Pa·s=1000 mPa·s",
        "1 Pa·s=1000 mP·s": "1 Pa·s=1000 mPa·s",
        "1P=1000cP": "1 P=100 cP",
        "0.006895kPa": "0.006895 MPa",
        "551.6kPa": "551.6 MPa",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = re.sub(r"(?<=\d)(?=(?:mm|cm|km|mPa|kPa|MPa|GPa|kg|g/cm³|kg/m³)\b)", " ", text)
    text = re.sub(r"(?<=\d)℃", " ℃", text)
    text = text.replace("mPa·S", "mPa·s").replace("Pa·S", "Pa·s")
    return text


def repair_markers(question_id: int, text: str) -> str:
    if question_id == 103:
        text = text.replace(
            "④]在键槽中遇阻、拉力稍大，转动钻柱很困难，但只要放下钻柱脱离键槽则旋转自如]",
            "④[在键槽中遇阻、拉力稍大，转动钻柱很困难，但只要放下钻柱脱离键槽则旋转自如]",
        )
    if question_id in (135, 240):
        text = text.replace("③]", "③[")
    if question_id == 135:
        text = text.replace("(3)]", "(3)[")
    if question_id == 39:
        text = text.replace("地层的]时代]", "地层的[时代]")
    if question_id == 40:
        text = text.replace("[提高采收率及开发速，用来进行", "[提高采收率及开发速度]，用来进行")
    if question_id == 355:
        text = text.replace("起到降温]的作用", "起到[降温]的作用")
    if question_id == 368:
        text = text.replace("[有效作用距离和裂缝的[导流能力]", "[有效作用距离]和裂缝的[导流能力]")
    return text


def split_long_clozes(text: str) -> str:
    pattern = re.compile(r"\[([^\[\]\r\n]{35,})\]")

    def split(match):
        answer = match.group(1)
        pieces = re.split(r"([，；。])", answer)
        output = []
        for index in range(0, len(pieces), 2):
            clause = pieces[index].strip()
            punctuation = pieces[index + 1] if index + 1 < len(pieces) else ""
            if clause:
                output.append(f"[{clause}]")
            output.append(punctuation)
        candidate = "".join(output)
        return candidate if candidate.count("[") > 1 else match.group(0)

    return pattern.sub(split, text)


def add_mnemonic_clozes(text: str) -> str:
    lines = []
    for line in text.split("\n"):
        stripped = line.strip()
        unwrapped = stripped[1:-1] if stripped.startswith("[") and stripped.endswith("]") else stripped
        if unwrapped in MNEMONICS:
            prefix = line[: len(line) - len(line.lstrip())]
            line = f"{prefix}[{unwrapped}]"
        lines.append(line)
    return "\n".join(lines)


def apply_question_repairs(question_id: int, text: str, *, manual: bool) -> str:
    if question_id == 4 and manual:
        text = text.replace(
            "每个[孔隙]所连通的[喉道数]。介于[2～15]之间。",
            "每个[孔隙]所[连通]的[喉道数]。介于[2～15]之间。",
        )
    if question_id == 9 and manual:
        text = text.replace(
            "流量Q与岩石渗透率K、截面积A和压差ΔP成正比，与流体黏度μ和砂体长度L成反比。",
            "流量 Q 与岩石[渗透率 K]、[截面积 A]和[压差 ΔP]成[正比]，"
            "与流体[黏度 μ]和砂体[长度 L]成[反比]。",
        )
    if question_id == 12 and manual:
        text = text.replace(
            "束缚水（残余水或不可动水）",
            "束缚水（[残余水或不可动水]）",
        )
    if question_id == 14 and manual:
        text = text.replace(
            "[滞留]或[闭锁]在岩石孔隙中",
            "[滞留]或[闭锁]在[岩石孔隙]中",
        )
    if question_id == 15 and manual:
        text = text.replace("未被[工作剂驱替]", "[未被工作剂驱替]")
    if question_id == 53:
        text = text.replace("俗称[泥浆]\n", "俗称[泥浆]。\n") if manual else text.replace("俗称泥浆\n", "俗称泥浆。\n")
    if question_id == 62:
        old = "钻井液密度(钻井液比重)：指单位体积钻井液的质量，单位g/cm³(或kg/m³)"
        if manual:
            new = "钻井液密度（钻井液比重）：指[单位体积钻井液的质量]，单位为[g/cm³]（或[kg/m³]）。"
        else:
            new = "钻井液密度（钻井液比重）：指单位体积钻井液的质量，单位为 g/cm³（或 kg/m³）。"
        text = text.replace(old, new)
    if question_id == 73:
        text = text.replace(
            "①[减轻对钻具的腐蚀]②[可预防氢脆引起钻具和套管损害]③[抑制钙镁盐溶解]④[充分发挥处理剂性能]。",
            "①[减轻对钻具的腐蚀]；②[预防氢脆引起钻具和套管损害]；"
            "③[抑制钙镁盐溶解]；④[充分发挥处理剂性能]。",
        )
    if question_id == 118 and manual:
        text = text.replace(
            "[钻下部地层采用重钻井液时产生的井内压力不致压裂上层套管处最薄弱的裸露地层]",
            "钻下部地层采用[重钻井液]时，产生的[井内压力]不致压裂"
            "[上层套管处最薄弱的裸露地层]",
        )
        text = text.replace(
            "[井内钻井液柱的压力和地层压力之间的压差不致产生压差卡套管现象]",
            "[井内钻井液柱压力与地层压力之间的压差]不致产生[压差卡套管]现象",
        )
    if question_id == 120:
        formula = "N-80=80 000 lbf/in²=80 000 psi=80kpsi=80000×0.006895 MPa=\n551.6 MPa"
        replacement = "[N-80=80 000 lbf/in²=80 ksi=551.6 MPa]" if manual else "N-80=80 000 lbf/in²=80 ksi=551.6 MPa"
        text = text.replace(formula, replacement)
        text = text.replace("单位为lbf/in²", "单位为 lbf/in²")
    if question_id == 127 and manual:
        text = text.replace(
            "只要保证套管柱上任意截面处的套管强度>外载，套管柱就是安全的。",
            "只要保证套管柱上[任意截面处的套管强度]>[外载]，套管柱就是安全的。",
        )
    if question_id == 202:
        text = text.replace("正循环洗井：洗井液油管", "正循环洗井：洗井液由油管")
    if question_id == 298:
        replacements = {
            "10⁸t": "10⁸ t",
            "5000×10⁴-1×10⁸t": "5000×10⁴～1×10⁸ t",
            "1000×10⁴~ 5000×10⁴t": "1000×10⁴～5000×10⁴ t",
            "500×10⁴~1000×10⁴t": "500×10⁴～1000×10⁴ t",
            "500×10⁴t": "500×10⁴ t",
        }
        for old, new in replacements.items():
            text = text.replace(old, new)
    if question_id == 319 and text.strip() == "采油工艺的分类\n[图片：/media/image32.jpeg]":
        if manual:
            text = (
                "采油工艺的分类\n"
                "采油工艺技术分为[自喷]和[人工举升]。\n"
                "人工举升分为[深井泵采油]和[气举采油]。\n"
                "深井泵采油分为[无杆泵采油]和[有杆泵采油]；无杆泵采油包括"
                "[电潜泵采油]、[射流泵采油]、[水力活塞泵采油]和[螺杆泵采油]，"
                "有杆泵采油包括[游梁式抽油机（主要）]和[无游梁式抽油机]。\n"
                "气举采油分为[连续气举]和[间歇气举]，间歇气举包括[腔式气举]和[柱塞气举]。\n"
                "[图片：/media/image32.jpeg]"
            )
        else:
            text = (
                "采油工艺的分类\n"
                "采油工艺技术分为自喷和人工举升。\n"
                "人工举升分为深井泵采油和气举采油。\n"
                "深井泵采油分为无杆泵采油和有杆泵采油；无杆泵采油包括"
                "电潜泵采油、射流泵采油、水力活塞泵采油和螺杆泵采油，"
                "有杆泵采油包括游梁式抽油机（主要）和无游梁式抽油机。\n"
                "气举采油分为连续气举和间歇气举，间歇气举包括腔式气举和柱塞气举。\n"
                "[图片：/media/image32.jpeg]"
            )
    if question_id == 325 and manual:
        text = text.replace(
            "[研究油井由于污染或采取增产措施后引起的完善性（或流动效率）改变所带来的影响]",
            "研究油井由于[污染]或[采取增产措施]后引起的"
            "[完善性（或流动效率）改变]所带来的影响",
        )
    return text


def clean_text(question_id: int, text: str, *, manual: bool) -> str:
    text = repair_markers(question_id, str(text or ""))
    for old, new in REPLACEMENTS.items():
        text = text.replace(old, new)
    text = repair_markers(question_id, text)
    text = normalize_units(normalize_punctuation(text))
    if question_id == 22:
        text = text.replace("[1 P（泊）=1000 cP]", "[1 P（泊）=100 cP]")
        text = text.replace("[1 Pa·s=1000 mP·s]", "[1 Pa·s=1000 mPa·s]")
    if question_id == 37:
        text = text.replace("C区——[水单相渗流区]、", "C区——[水单相渗流区]。")
    if question_id == 54:
        text = text.replace("[无固相钻井液]：", "[无固相钻井液]。")
    if question_id == 105:
        text = text.replace("②[泡油解卡]；向井内", "②[泡油解卡]：向井内")
    if question_id == 120:
        text = text.replace("80000lbf/in²", "80 000 lbf/in²").replace("80000 psi", "80 000 psi")
    if question_id == 295:
        text = text.replace("5000×10⁴-1×10⁸t", "5000×10⁴～1×10⁸ t")
    text = apply_question_repairs(question_id, text, manual=manual)
    if manual:
        text = text.replace("[(性蜜罐漏拔稳)]", "[性蜜罐漏拔稳]")
        text = add_mnemonic_clozes(split_long_clozes(text))
    return text


def strip_markers(text: str) -> str:
    return re.sub(r"[\[\]]", "", text).strip().rstrip("：:")


def restructure_manual(question_id: int, prompt: str, manual: str) -> tuple[str, str]:
    prompt = TITLE_OVERRIDES.get(question_id, strip_markers(prompt))
    prompt = normalize_punctuation(prompt).strip().rstrip("。；，")
    lines = manual.strip("\n").split("\n") if manual else []
    if not lines:
        return prompt, prompt

    if question_id == 18 and lines[0].startswith("[图片："):
        lines.insert(0, prompt)
    elif question_id == 19:
        lines[0] = "地层油的体积系数 Bo（又称[原油地下体积系数]）"
        if len(lines) > 1 and not lines[1].lstrip().startswith("是指"):
            lines[1] = "是指" + lines[1]
    elif question_id in ANSWER_BEARING_TITLES:
        if strip_markers(lines[0]) != prompt:
            lines.insert(0, prompt)
        else:
            lines[0] = prompt
    elif question_id in TITLE_ONLY_ROWS:
        lines[0] = prompt
    else:
        first_plain = strip_markers(lines[0])
        if first_plain.rstrip("。") == prompt.rstrip("。"):
            lines[0] = prompt
        elif lines[0].startswith(prompt + "：") or lines[0].startswith(prompt + ":"):
            _, body = re.split(r"[：:]", lines[0], maxsplit=1)
            lines[0:1] = [prompt, body.strip()]
        elif lines[0].startswith("[图片："):
            lines.insert(0, prompt)

    # Prompts are questions/labels and never end in a colon.
    lines[0] = re.sub(r"[：:]\s*$", "", lines[0].strip())
    return prompt, "\n".join(lines)


def infer_types(title: str, manual: str) -> list[str]:
    combined = title + "\n" + manual
    types = []
    rules = [
        ("definition", r"定义|概念|什么是|是指|称为"),
        ("classification", r"分类|类型|分为|包括"),
        ("condition", r"条件|适用|要求|原则|前提"),
        ("mechanism", r"机理|机制|过程"),
        ("effect", r"作用|意义|目的|功能|优点|缺点|影响|措施|预防|处理"),
        ("result", r"结果|后果|特征|规律|关系"),
        ("formula", r"公式|定律|＝|=|/|系数"),
        ("value", r"\d+(?:\.\d+)?\s*(?:%|℃|mm|cm|m|Pa|MPa|mPa|kg|g/cm³)"),
        ("example", r"例如|如：|举例"),
    ]
    for category, pattern in rules:
        if re.search(pattern, combined):
            types.append(category)
    if any(mnemonic in combined for mnemonic in MNEMONICS):
        types.append("mnemonic")
    if re.search(r"(?m)^(?:①|（1）|\*)", manual):
        types.append("list")
    return types or ["definition"]


def ensure_column(sheet, header: str) -> int:
    for cell in sheet[1]:
        if cell.value == header:
            return cell.column
    column = sheet.max_column + 1
    source = sheet.cell(1, 1)
    target = sheet.cell(1, column, header)
    target.fill = copy(source.fill)
    target.font = copy(source.font)
    target.alignment = copy(source.alignment)
    return column


def update_manual_sheet(workbook, changes: list[tuple]):
    sheet = workbook[MANUAL_SHEET]
    headers = {cell.value: cell.column for cell in sheet[1]}
    type_column = ensure_column(sheet, "知识类型")
    sheet.column_dimensions[sheet.cell(1, type_column).column_letter].width = 25

    for row in range(2, sheet.max_row + 1):
        question_id = int(sheet.cell(row, headers["题号"]).value)
        prompt_cell = sheet.cell(row, headers["题目提示"])
        original_cell = sheet.cell(row, headers["题目原文（勿改）"])
        manual_cell = sheet.cell(row, headers["手动挖空版（在此编辑）"])
        before = str(manual_cell.value or "")
        original = clean_text(question_id, original_cell.value, manual=False)
        manual = clean_text(question_id, manual_cell.value, manual=True)
        prompt, manual = restructure_manual(question_id, prompt_cell.value, manual)
        prompt_cell.value = prompt
        original_cell.value = original
        manual_cell.value = manual
        types = infer_types(prompt, manual)
        sheet.cell(row, type_column, "、".join(QUESTION_TYPES[item] for item in types))
        sheet.cell(row, type_column).alignment = Alignment(wrap_text=True, vertical="top")
        if before != manual:
            changes.append((question_id, "正文、格式、挖空", "已按 1.1.0 题库规范校正"))


def update_detail_sheet(workbook):
    sheet = workbook["块级明细"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    type_column = ensure_column(sheet, "知识类型")
    sheet.column_dimensions[sheet.cell(1, type_column).column_letter].width = 18
    for row in range(2, sheet.max_row + 1):
        value = sheet.cell(row, headers["题号"]).value
        if value is None:
            continue
        question_id = int(value)
        original_cell = sheet.cell(row, headers["原始内容（勿改）"])
        manual_cell = sheet.cell(row, headers["手动挖空内容"])
        if original_cell.value:
            original_cell.value = clean_text(question_id, original_cell.value, manual=False)
        if manual_cell.value:
            manual_cell.value = clean_text(question_id, manual_cell.value, manual=True)
        text = str(manual_cell.value or "")
        types = infer_types(text, text)
        sheet.cell(row, type_column, "、".join(QUESTION_TYPES[item] for item in types))
        sheet.cell(row, type_column).alignment = Alignment(wrap_text=True, vertical="top")


def update_guide(workbook):
    sheet = workbook["使用说明"]
    updates = {
        "题库规模": "369 道当前有效题目；题号连续排列，sourceId 保留原始题号以兼容存档。",
        "当前状态": "本文件是 1.1.0 的人工挖空权威源，已完成 OCR、单位、标题与口诀空审校。",
        "标记方法": "把需要隐藏的完整答案放在半角方括号中，例如：[孔隙]与[喉道][直径]的比值。",
        "语义要求": "短空保持原设计；过长答案按条件、动作、原因和结果拆分。定义、作用、意义、分类等提示词保持可见。",
        "文字纠错": "题目原文与手动挖空版均按教材核对；公式、单位和特殊符号遵循科学写法。",
        "后续导入": "App 通过 tools/import_manual_questions.py 读取此文件；标题空、口诀空、分类和 sourceId 均会保留。",
    }
    for row in range(1, sheet.max_row + 1):
        label = sheet.cell(row, 1).value
        if label in updates:
            sheet.cell(row, 2, updates[label])


def add_audit_sheet(workbook, changes):
    name = "1.1.0审校记录"
    if name in workbook.sheetnames:
        sheet = workbook[name]
        sheet.delete_rows(1, sheet.max_row)
    else:
        sheet = workbook.create_sheet(name)
    sheet.append(["题号", "审校范围", "处理结果"])
    for item in changes:
        sheet.append(item)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="176B5B")
        cell.font = Font(color="FFFFFF", bold=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 10
    sheet.column_dimensions["B"].width = 28
    sheet.column_dimensions["C"].width = 62
    sheet["E1"] = "审校时间"
    sheet["F1"] = datetime.now().strftime("%Y-%m-%d %H:%M")


def validate(workbook):
    sheet = workbook[MANUAL_SHEET]
    headers = {cell.value: cell.column for cell in sheet[1]}
    ids = []
    problems = []
    mnemonic_count = 0
    for row in range(2, sheet.max_row + 1):
        question_id = int(sheet.cell(row, headers["题号"]).value)
        ids.append(question_id)
        prompt = str(sheet.cell(row, headers["题目提示"]).value or "")
        manual = str(sheet.cell(row, headers["手动挖空版（在此编辑）"]).value or "")
        if prompt.endswith(("：", ":")):
            problems.append(f"第 {question_id} 题提示仍以冒号结尾")
        if re.search(r"[\ue000-\uf8ff]", manual):
            problems.append(f"第 {question_id} 题仍含私有区字符")
        if manual.count("[") != manual.count("]"):
            problems.append(f"第 {question_id} 题方括号不平衡")
        for mnemonic in MNEMONICS:
            if mnemonic in manual:
                if f"[{mnemonic}]" not in manual:
                    problems.append(f"第 {question_id} 题口诀未挖空：{mnemonic}")
                mnemonic_count += 1
        if any(token in manual for token in ("完并", "井简", "气审井", "1000cP", "mP·s", "有效作用距和")):
            problems.append(f"第 {question_id} 题仍含已知 OCR/单位错误")
    if len(ids) != len(set(ids)):
        problems.append("题号重复")
    if problems:
        raise ValueError("\n".join(problems[:30]))
    return {"rows": len(ids), "mnemonics": mnemonic_count}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", nargs="?", type=Path, default=DEFAULT_WORKBOOK)
    args = parser.parse_args()
    workbook = load_workbook(args.workbook)
    changes = []
    update_manual_sheet(workbook, changes)
    update_detail_sheet(workbook)
    update_guide(workbook)
    add_audit_sheet(workbook, changes)
    result = validate(workbook)
    workbook.save(args.workbook)
    print({"workbook": str(args.workbook), **result, "changed": len(changes)})


if __name__ == "__main__":
    main()
