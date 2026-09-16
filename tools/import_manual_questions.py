import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "public" / "data" / "questions.json"
XLSX_PATH = ROOT / "829题库.xlsx"

IMAGE_RE = re.compile(r"^\[图片：([^\]]+)\]$")
MARK_RE = re.compile(r"\[([^\[\]\r\n]+)\]")
PRIVATE_RE = re.compile(r"[\ue000-\uf8ff]")

MNEMONICS = {
    "性蜜罐漏拔稳",
    "高侵交（胶）体（替）毛（冒）",
    "膨固破沉锁",
    "密桥类完活",
    "热砂通渗侧隔低裂",
    "安稳护衡调配眼影粒腐",
    "速步术式原层",
    "地位独规",
    "术能环稳效速率",
    "天性关压",
    "基发适提",
    "厚隔技物",
    "则透后速造",
    "速调一面",
    "整密稀换局",
    "要原储层工技输流经地温驱指（可数）",
    "描藏油地经综",
    "程试监预",
    "分划同适剔细",
    "强粒杂便密",
    "地能转错情",
    "悬砂滤网摩擦稳定，这是低配的经济。",
}

TYPE_NAMES = {
    "定义": "definition",
    "分类": "classification",
    "条件": "condition",
    "机理": "mechanism",
    "作用": "effect",
    "结果": "result",
    "公式": "formula",
    "数值": "value",
    "口诀": "mnemonic",
    "条目": "list",
    "例子": "example",
}


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # Normalize full-width spaces and accidental OCR spacing without touching
    # intentional spaces in formulas and units.
    text = text.replace("\u00a0", " ").replace("\u3000", " ")
    return text


def title_without_colon(title: str) -> str:
    return re.sub(r"[：:]\s*$", "", title.strip())


def cloze_type(answer: str) -> str:
    if answer in MNEMONICS:
        return "mnemonic"
    if re.search(r"[=<>≤≥/]", answer) or re.search(r"\b(?:Swi|Sor|pwf|Bo|Bg|Rs|IPR)\b", answer):
        return "formula"
    if re.search(r"\d", answer):
        return "value"
    return "keyword"


def paragraph_category(text: str, question_types: list[str]) -> str:
    stripped = text.strip()
    if stripped in MNEMONICS:
        return "mnemonic"
    rules = [
        ("definition", r"^(?:定义|概念)|是指|称为"),
        ("classification", r"分类|类型|分为|包括"),
        ("condition", r"条件|要求|原则|适用"),
        ("mechanism", r"机理|机制|过程"),
        ("effect", r"作用|意义|目的|功能|优点|缺点|措施|预防|处理"),
        ("result", r"结果|后果|特征|规律|关系"),
        ("formula", r"公式|定律|[=＝]"),
        ("value", r"\d+(?:\.\d+)?\s*(?:%|℃|mm|cm|m|Pa|MPa|mPa|kg|g/cm³)"),
    ]
    for category, pattern in rules:
        if re.search(pattern, stripped):
            return category
    return question_types[0] if question_types else "list"


def parse_line(line: str, question_types: list[str]):
    line = normalize_text(line)
    image_match = IMAGE_RE.fullmatch(line.strip())
    if image_match:
        return {"type": "image", "src": image_match.group(1).strip()}

    text_parts = []
    clozes = []
    cursor = 0
    for match in MARK_RE.finditer(line):
        text_parts.append(line[cursor:match.start()])
        answer = match.group(1)
        start = sum(len(part) for part in text_parts)
        text_parts.append(answer)
        end = start + len(answer)
        clozes.append({
            "start": start,
            "end": end,
            "answer": answer,
            "clozeType": cloze_type(answer),
        })
        cursor = match.end()
    text_parts.append(line[cursor:])
    text = "".join(text_parts)
    return {
        "type": "paragraph",
        "text": text,
        "category": paragraph_category(text, question_types),
        "clozes": clozes,
    }


def parse_manual(question_id: int, manual: str, question_types: list[str]):
    manual = normalize_text(manual or "").strip("\n")
    if not manual or manual in {"删去", "删除", "重复题", "暂不处理"}:
        return None

    blocks = []
    for line in manual.split("\n"):
        if line == "":
            if blocks and blocks[-1].get("type") != "spacer":
                blocks.append({"type": "spacer"})
            continue
        blocks.append(parse_line(line, question_types))

    while blocks and blocks[-1].get("type") == "spacer":
        blocks.pop()
    return blocks


def expected_cloze_count(manual: str) -> int:
    total = 0
    for line in normalize_text(manual or "").split("\n"):
        if IMAGE_RE.fullmatch(line.strip()):
            continue
        total += len(MARK_RE.findall(line))
    return total


def read_manual_rows():
    workbook = load_workbook(XLSX_PATH, read_only=True, data_only=False)
    sheet = workbook["逐题人工挖空"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    indexes = {name: index for index, name in enumerate(headers)}
    rows = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        value = row[indexes["题号"]]
        if value is None:
            continue
        question_id = int(value)
        rows[question_id] = {
            "manual": row[indexes["手动挖空版（在此编辑）"]] or "",
            "status": row[indexes["审核状态"]] or "",
            "prompt": row[indexes["题目提示"]] or "",
            "chapterNumber": int(row[indexes["章"]] or 0),
            "chapterTitle": row[indexes["章节名称"]] or "",
            "knowledgeTypes": [
                TYPE_NAMES[item]
                for item in str(row[indexes.get("知识类型", -1)] or "").split("、")
                if item in TYPE_NAMES
            ] if "知识类型" in indexes else [],
        }
    return rows


def validate_question(question):
    if not question["blocks"]:
        raise ValueError(f"第 {question['id']} 题没有内容")
    if PRIVATE_RE.search(json.dumps(question, ensure_ascii=False)):
        raise ValueError(f"第 {question['id']} 题仍含私有字符")
    for block in question["blocks"]:
        if block["type"] != "paragraph":
            continue
        previous_end = -1
        for cloze in block.get("clozes", []):
            start, end = cloze["start"], cloze["end"]
            if not (isinstance(start, int) and isinstance(end, int) and 0 <= start < end):
                raise ValueError(f"第 {question['id']} 题存在无效挖空范围")
            if start < previous_end or block["text"][start:end] != cloze["answer"]:
                raise ValueError(f"第 {question['id']} 题挖空范围重叠或答案不一致")
            previous_end = end
        for marker in ("[", "]"):
            if marker in block["text"]:
                raise ValueError(f"第 {question['id']} 题正文残留未解析标记: {block['text']}")


def main():
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    manual_rows = read_manual_rows()
    original = {
        question.get("sourceId", question["id"]): question
        for question in data["questions"]
    }
    questions = []
    removed = []

    for question_id in sorted(manual_rows):
        row = manual_rows[question_id]
        blocks = parse_manual(question_id, row["manual"], row["knowledgeTypes"])
        if blocks is None:
            removed.append(question_id)
            continue
        question = original.get(question_id, {"id": question_id})
        question["id"] = question_id
        question["sourceId"] = question_id
        question["blocks"] = blocks
        question["title"] = title_without_colon(row["prompt"] or f"第 {question_id} 题")
        question["knowledgeTypes"] = row["knowledgeTypes"]
        question["category"] = row["knowledgeTypes"][0] if row["knowledgeTypes"] else "list"
        question["chapter"] = {
            "number": row["chapterNumber"],
            "title": row["chapterTitle"],
        }
        question["fillCount"] = sum(
            len(block.get("clozes", [])) for block in blocks if block["type"] == "paragraph"
        )
        expected = expected_cloze_count(row["manual"])
        if question["fillCount"] != expected:
            raise ValueError(
                f"第 {question_id} 题 Excel 有 {expected} 个空，导入后为 {question['fillCount']} 个"
            )
        validate_question(question)
        questions.append(question)

    # The workbook itself is the source of truth for deletions.  Ensure an
    # accidental omission cannot silently drop a question from the app.
    expected_ids = set(original) - set(removed)
    actual_ids = {question["id"] for question in questions}
    missing = expected_ids - actual_ids
    if missing:
        raise ValueError(f"Excel 未提供题目: {sorted(missing)}")

    # App-facing numbers stay continuous after workbook rows are deleted.
    # sourceId remains stable so future textbook audits can still address the
    # original Excel row without coupling saved progress to that old number.
    for display_id, question in enumerate(questions, start=1):
        question["sourceId"] = question["id"]
        question["id"] = display_id

    data["version"] = 6
    data["source"] = XLSX_PATH.name
    data["schema"] = "semantic-cloze-v2"
    data["total"] = len(questions)
    data["questions"] = questions
    DATA_PATH.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(json.dumps({
        "version": data["version"],
        "questions": len(questions),
        "removed": removed,
        "clozes": sum(question["fillCount"] for question in questions),
        "mnemonics": sum(
            cloze.get("clozeType") == "mnemonic"
            for question in questions
            for block in question["blocks"] if block["type"] == "paragraph"
            for cloze in block.get("clozes", [])
        ),
        "nofill": [question["id"] for question in questions if not question["fillCount"]],
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
