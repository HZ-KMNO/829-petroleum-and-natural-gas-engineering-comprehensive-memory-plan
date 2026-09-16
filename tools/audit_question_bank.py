import json
import re
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "public" / "data" / "questions.json"
XLSX_PATH = ROOT / "829题库.xlsx"
MARK_RE = re.compile(r"\[([^\[\]\r\n]+)\]")
IMAGE_RE = re.compile(r"^\[图片：([^\]]+)\]$")
PRIVATE_RE = re.compile(r"[\ue000-\uf8ff]")
INVALID_TEXT = (
    "完并", "井简", "气审井", "携带赃物", "配置高密度钻井液",
    "1000cP", "mP·s", "CaMg(CO)3", "油单向流动区",
)


def workbook_rows():
    workbook = load_workbook(XLSX_PATH, read_only=True, data_only=False)
    sheet = workbook["逐题人工挖空"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    indexes = {name: index for index, name in enumerate(headers)}
    rows = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        source_id = int(row[indexes["题号"]])
        manual = str(row[indexes["手动挖空版（在此编辑）"]] or "")
        if manual.strip() in {"删去", "删除", "重复题", "暂不处理"}:
            continue
        expected = sum(
            len(MARK_RE.findall(line))
            for line in manual.splitlines()
            if not IMAGE_RE.fullmatch(line.strip())
        )
        rows[source_id] = {
            "expectedClozes": expected,
            "prompt": str(row[indexes["题目提示"]] or ""),
        }
    return rows


def main():
    data = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    rows = workbook_rows()
    questions = data["questions"]
    errors = []
    media = set()
    mnemonic_count = 0
    cloze_count = 0

    if data.get("version") != 6 or data.get("schema") != "semantic-cloze-v2":
        errors.append("题库结构版本不是 v6 semantic-cloze-v2")
    if [question["id"] for question in questions] != list(range(1, len(questions) + 1)):
        errors.append("应用题号不连续")
    if set(rows) != {question["sourceId"] for question in questions}:
        errors.append("Excel 有效题目与 JSON sourceId 不一致")

    for question in questions:
        source_id = question["sourceId"]
        if "：" in question["title"] or ":" in question["title"]:
            errors.append(f"源题 {source_id} 标题仍含冒号")
        if not question.get("knowledgeTypes") or not question.get("category"):
            errors.append(f"源题 {source_id} 缺少知识类型")
        actual = 0
        for block in question["blocks"]:
            if block["type"] == "image":
                media.add(block["src"])
                continue
            if block["type"] != "paragraph":
                continue
            if not block.get("category"):
                errors.append(f"源题 {source_id} 正文块缺少类别")
            previous_end = -1
            for cloze in block.get("clozes", []):
                actual += 1
                cloze_count += 1
                start, end = cloze["start"], cloze["end"]
                if start < previous_end or block["text"][start:end] != cloze["answer"]:
                    errors.append(f"源题 {source_id} 挖空范围错误：{cloze['answer']}")
                if cloze.get("clozeType") not in {"keyword", "formula", "value", "mnemonic"}:
                    errors.append(f"源题 {source_id} 挖空缺少语义类型")
                if cloze.get("clozeType") == "mnemonic":
                    mnemonic_count += 1
                previous_end = end
        if actual != rows[source_id]["expectedClozes"] or actual != question["fillCount"]:
            errors.append(
                f"源题 {source_id} Excel/JSON 空数不一致："
                f"{rows[source_id]['expectedClozes']}/{actual}/{question['fillCount']}"
            )
        if actual == 0:
            errors.append(f"源题 {source_id} 没有挖空")

    serialized = json.dumps(data, ensure_ascii=False)
    if PRIVATE_RE.search(serialized):
        errors.append("JSON 仍含私用区乱码")
    for token in INVALID_TEXT:
        if token in serialized:
            errors.append(f"JSON 仍含已知错误：{token}")
    for source in media:
        if not (ROOT / "public" / source.lstrip("/")).is_file():
            errors.append(f"缺少图片：{source}")
    if mnemonic_count != 22:
        errors.append(f"口诀空应为 22，实际为 {mnemonic_count}")

    if errors:
        raise SystemExit("\n".join(errors[:50]))
    print(json.dumps({
        "version": data["version"],
        "questions": len(questions),
        "clozes": cloze_count,
        "mnemonics": mnemonic_count,
        "images": len(media),
        "categories": sorted({item for question in questions for item in question["knowledgeTypes"]}),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
